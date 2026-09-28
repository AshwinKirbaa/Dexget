import os
from typing import Optional, List

class DocumentExtractor:
    @staticmethod
    def extract_text(file_path: str) -> str:
        ext = file_path.lower().split('.')[-1]
        try:
            if ext in ["txt", "md", "csv"]:
                with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                    return f.read()
            elif ext == "pdf":
                try:
                    import PyPDF2
                    text = ""
                    with open(file_path, 'rb') as f:
                        reader = PyPDF2.PdfReader(f)
                        for page in reader.pages:
                            text += page.extract_text() + "\n"
                    return text
                except ImportError:
                    return "Error: PyPDF2 not installed."
            elif ext == "docx":
                try:
                    import docx
                    doc = docx.Document(file_path)
                    return "\n".join([p.text for p in doc.paragraphs])
                except ImportError:
                    return "Error: python-docx not installed."
            elif ext in ["png", "jpg", "jpeg", "bmp"]:
                from jarvis.memory.ocr import OCRProvider
                return OCRProvider.perform_ocr(file_path)
            elif ext == "xlsx":
                try:
                    import openpyxl
                    wb = openpyxl.load_workbook(file_path, data_only=True)
                    text = []
                    for sheet in wb.worksheets:
                        text.append(f"Sheet: {sheet.title}")
                        for row in sheet.iter_rows(values_only=True):
                            text.append("\t".join([str(c) if c is not None else "" for c in row]))
                    return "\n".join(text)
                except ImportError:
                    return "Error: openpyxl not installed."
            elif ext == "pptx":
                try:
                    from pptx import Presentation
                    prs = Presentation(file_path)
                    text = []
                    for slide in prs.slides:
                        for shape in slide.shapes:
                            if hasattr(shape, "text"):
                                text.append(shape.text)
                    return "\n".join(text)
                except ImportError:
                    return "Error: python-pptx not installed."
            else:
                return "Unsupported file format."
        except Exception as e:
            return f"Error extracting text: {e}"
